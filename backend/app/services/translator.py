import json
import re
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import requests
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from app.schemas import TranslationOptions


def parse_batch_json_robust(raw: str, expected_n: int) -> list[str] | None:
    try:
        data = json.loads(raw)
        arr = data.get("translations", [])
        if isinstance(arr, list) and len(arr) == expected_n:
            return [str(x.get("text", "")) if isinstance(x, dict) else str(x) for x in arr]
    except Exception:
        pass

    fenced = re.search(r"\{[\s\S]*\}", raw)
    if fenced:
        seg = fenced.group(0)
        try:
            data = json.loads(seg)
            for key in ("translations", "items", "results"):
                arr = data.get(key, [])
                if isinstance(arr, list) and len(arr) == expected_n:
                    out: list[str] = []
                    for item in arr:
                        if isinstance(item, dict):
                            out.append(str(item.get("text", item.get("translation", ""))))
                        else:
                            out.append(str(item))
                    return out
        except Exception:
            pass
    return None


@dataclass
class TextNode:
    slide: int
    kind: str
    paragraph: Any
    text: str


class PptTranslatorService:
    def __init__(self, options: TranslationOptions):
        self.options = options

    def translate_pptx(self, input_path: str | Path, output_path: str | Path) -> None:
        prs = Presentation(str(input_path))
        nodes = self._collect_text_nodes(prs)
        if not nodes:
            prs.save(str(output_path))
            return

        unique_items: dict[str, dict[str, str | int]] = {}
        for node in nodes:
            src = self._sanitize_text_for_ppt(node.text)
            key = f"{self._normalize_key(src)}||{node.kind}"
            if key not in unique_items:
                unique_items[key] = {"text": src, "kind": node.kind, "slide": node.slide}

        items = list(unique_items.values())
        batches = [
            items[i : i + self.options.batch_size]
            for i in range(0, len(items), self.options.batch_size)
        ]

        cache: dict[str, str] = {}
        with ThreadPoolExecutor(max_workers=self.options.max_workers) as executor:
            future_map = {executor.submit(self._translate_items, batch): batch for batch in batches}
            for future in as_completed(future_map):
                batch = future_map[future]
                try:
                    translated = future.result()
                    for src_item, tgt in zip(batch, translated):
                        key = f"{self._normalize_key(str(src_item['text']))}||{src_item['kind']}"
                        cache[key] = self._sanitize_text_for_ppt(str(tgt))
                except Exception:
                    for src_item in batch:
                        key = f"{self._normalize_key(str(src_item['text']))}||{src_item['kind']}"
                        cache[key] = str(src_item["text"])

        for node in nodes:
            src = self._sanitize_text_for_ppt(node.text)
            key = f"{self._normalize_key(src)}||{node.kind}"
            translated_text = cache.get(key, src)
            self._safe_replace_runs_text(node.paragraph, translated_text)

        prs.save(str(output_path))

    def _translate_items(self, items: list[dict[str, str | int]]) -> list[str]:
        prompt = self._build_prompt_batch(items)
        raw = self._call_chat(prompt)
        parsed = parse_batch_json_robust(raw, len(items))
        if parsed is not None:
            return [self._post_process_item(str(it["text"]), t) for it, t in zip(items, parsed)]

        out: list[str] = []
        for item in items:
            single_prompt = self._build_prompt_single(str(item["text"]))
            translated = self._call_chat(single_prompt)
            out.append(self._post_process_item(str(item["text"]), translated))
        return out

    def _post_process_item(self, src_text: str, translated: str) -> str:
        text = self._sanitize_text_for_ppt(translated)
        if self.options.force_terminology_replace:
            text = self._force_terminology_replace(text)
        text = self._sync_trailing_punct(src_text, text)
        return text

    def _call_chat(self, prompt: str) -> str:
        headers = {"Content-Type": "application/json"}
        if self.options.api_key.strip():
            headers["Authorization"] = f"Bearer {self.options.api_key.strip()}"

        body = {
            "model": self.options.model.strip() or "deepseek-v4-flash",
            "messages": [
                {"role": "system", "content": "You are a translation assistant. Follow constraints strictly."},
                {"role": "user", "content": prompt},
            ],
            "chat_template_kwargs": {"thinking": False},
        }

        last_error: Exception | None = None
        for attempt in range(max(1, self.options.retry)):
            try:
                response = requests.post(
                    self.options.endpoint.strip(),
                    headers=headers,
                    data=json.dumps(body),
                    timeout=240,
                )
                if response.status_code != 200:
                    raise RuntimeError(f"HTTP {response.status_code}: {response.text[:500]}")
                data = response.json()
                return data["choices"][0]["message"]["content"].strip()
            except Exception as exc:
                last_error = exc
                if attempt < self.options.retry - 1:
                    time.sleep(float(attempt + 1))
                else:
                    raise RuntimeError(f"API failed after retries: {last_error}") from exc
        raise RuntimeError(f"API failed after retries: {last_error}")

    def _build_prompt_batch(self, items: list[dict[str, str | int]]) -> str:
        terminology = "No terminology provided."
        if self.options.terminology:
            terminology = "\n".join(f"- {k} => {v}" for k, v in self.options.terminology.items())

        style_hint = ""
        if self.options.translation_style.lower() == "minimal":
            style_hint = "Use concise business-slide style and keep only core message."
        elif self.options.translation_style.lower() == "balanced":
            style_hint = "Keep full meaning but concise business-slide wording."

        payload = [{"text": str(item["text"]), "kind": str(item["kind"])} for item in items]
        return (
            f"You are a professional translator.\n"
            f"Translate from {self.options.source_language} to {self.options.target_language}.\n"
            f"{style_hint}\n\n"
            f"Terminology:\n{terminology}\n\n"
            f"STRICT OUTPUT FORMAT:\n"
            f'{{"translations":[{{"text":"..."}},{{"text":"..."}}]}}\n'
            f"Count/order must match input exactly.\n"
            f"No markdown, no commentary.\n\n"
            f"Input:\n{json.dumps(payload, ensure_ascii=False)}"
        )

    def _build_prompt_single(self, text: str) -> str:
        terminology = "No terminology provided."
        if self.options.terminology:
            terminology = "\n".join(f"- {k} => {v}" for k, v in self.options.terminology.items())
        template = self.options.prompt_template.strip()
        try:
            base = template.format(
                source_language=self.options.source_language,
                target_language=self.options.target_language,
                terminology=terminology,
                text=text,
            )
        except Exception:
            base = (
                f"Translate from {self.options.source_language} to {self.options.target_language}.\n"
                f"Terminology:\n{terminology}\n\nText:\n{text}"
            )
        return f"{base}\n\nReturn translation only."

    def _iter_text_frames_in_shape(self, shape: Any):
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for sub in shape.shapes:
                yield from self._iter_text_frames_in_shape(sub)
            return
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame and shape.text_frame:
            yield shape.text_frame, shape
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    if cell.text_frame:
                        yield cell.text_frame, shape

    def _collect_text_nodes(self, prs: Presentation) -> list[TextNode]:
        nodes: list[TextNode] = []
        for slide_index, slide in enumerate(prs.slides, start=1):
            for shape in slide.shapes:
                for text_frame, owner in self._iter_text_frames_in_shape(shape):
                    for paragraph in text_frame.paragraphs:
                        txt = (paragraph.text or "").strip()
                        if txt:
                            nodes.append(
                                TextNode(
                                    slide=slide_index,
                                    kind=self._detect_para_kind(owner, paragraph),
                                    paragraph=paragraph,
                                    text=paragraph.text,
                                )
                            )
        return nodes

    def _detect_para_kind(self, shape: Any, para: Any) -> str:
        try:
            if getattr(shape, "is_placeholder", False):
                if shape.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
                    return "title"
        except Exception:
            pass
        try:
            for run in para.runs:
                if run.font and run.font.size and run.font.size.pt >= 24:
                    return "title"
        except Exception:
            pass
        return "body"

    def _safe_replace_runs_text(self, para: Any, new_text: str) -> None:
        runs = list(para.runs)
        if not runs:
            para.text = new_text
            return
        target_run = max(runs, key=lambda run: len((run.text or "").strip()))
        target_run.text = new_text or ""
        for run in runs:
            if run is not target_run:
                run.text = ""

    def _sanitize_text_for_ppt(self, text: str) -> str:
        if text is None:
            return ""
        out = text.replace("\r\n", "\n").replace("\r", "\n").strip()
        out = re.sub(r"\n{3,}", "\n\n", out)
        return out

    def _normalize_key(self, text: str) -> str:
        return re.sub(r"\s+", "", text or "").lower()

    def _term_regex_pattern(self, term: str):
        esc = re.escape(term.strip())
        if re.search(r"[A-Za-z]", term):
            return re.compile(rf"\b{esc}\b", re.I)
        return re.compile(esc)

    def _force_terminology_replace(self, text: str) -> str:
        out = text
        for src, tgt in self.options.terminology.items():
            src = (src or "").strip()
            tgt = (tgt or "").strip()
            if not src:
                continue
            pattern = self._term_regex_pattern(src)
            out = pattern.sub(tgt, out)
        return out

    def _sync_trailing_punct(self, src_text: str, tgt_text: str) -> str:
        src = (src_text or "").rstrip()
        tgt = (tgt_text or "").rstrip()
        if not src or not tgt:
            return tgt_text
        sentence_punct = ".!?;:。！？；："
        src_has = src[-1] in sentence_punct
        tgt_has = tgt[-1] in sentence_punct
        if src_has and not tgt_has:
            return tgt + src[-1]
        if not src_has and tgt_has:
            return tgt[:-1]
        return tgt

